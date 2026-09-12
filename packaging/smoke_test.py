"""Proves a DocShift build converts real files, both ways, as a user would.

    python packaging/smoke_test.py dist/DocShift.exe
    python packaging/smoke_test.py python -m docshift

CI runs it against both: the source, and the .exe the Windows job just built.

The PDF it makes has text on two pages, an image, and a curve and a diagonal
line. The last two matter most. pdf2docx finds drawings like those with OpenCV
and clips them to a bitmap, so a build missing part of OpenCV fails here --
rather than on someone's first real document.

The Word document it makes has a heading, a list, a table and a picture, which
is what proves the other direction: python-docx reading it needs the template
data a frozen build can easily leave out, and drawing it needs MuPDF's story
engine.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pymupdf


def make_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Shifted by DocShift", fontsize=14)
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
    image.set_rect(image.irect, (220, 40, 60))
    page.insert_image(pymupdf.Rect(72, 100, 200, 228), pixmap=image)
    page.draw_circle((360, 300), 40, color=(0.1, 0.1, 0.1), fill=(0.9, 0.2, 0.3))
    page.draw_line((72, 400), (300, 500), color=(0, 0, 0), width=2)
    document.new_page().insert_text((72, 72), "The second page survives", fontsize=14)
    document.save(str(path))
    document.close()


def make_docx(path: Path) -> None:
    from docx.shared import Inches

    document = docx.Document()
    document.add_heading("Shifted by DocShift", level=1)
    document.add_paragraph("The second paragraph survives")
    document.add_paragraph("A bullet that survives", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "New York"

    picture = path.with_name("picture.png")
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
    image.set_rect(image.irect, (40, 120, 220))
    image.save(str(picture))
    document.add_picture(str(picture), width=Inches(1.5))
    picture.unlink()

    document.save(str(path))


def run(command: list[str], *args: object) -> int:
    return subprocess.run([*command, *map(str, args)], timeout=300).returncode


def fail(message: str) -> None:
    # A GitHub Actions error annotation; plain enough to read anywhere else.
    print(f"::error title=Smoke test failed::{message}")
    sys.exit(1)


def main(command: list[str]) -> int:
    if Path(command[0]).is_file():
        command = [str(Path(command[0]).resolve()), *command[1:]]

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        pdf = work / "sample.pdf"
        make_pdf(pdf)

        code = run(command, pdf, "--output", work / "out")
        if code != 0:
            fail(f"converting a real PDF exited with {code}")
        result = work / "out" / "sample.docx"
        if not result.exists():
            fail("it exited cleanly but wrote no DOCX")
        document = docx.Document(str(result))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for line in ("Shifted by DocShift", "The second page survives"):
            if line not in text:
                fail(f"{line!r} is missing from the DOCX")
        images = sum("image" in rel.reltype for rel in document.part.rels.values())
        if images < 2:
            fail(f"expected the image plus the drawings OpenCV clips, found {images} image(s)")
        print(f"converted: text from both pages, and {images} images")

        code = run(command, pdf, "--output", work / "out")
        if code != 0 or not (work / "out" / "sample (1).docx").exists():
            fail("a second run did not save beside the first as 'sample (1).docx'")
        print("a second run saved as 'sample (1).docx' instead of overwriting")

        # Every run passes --output. Without a console, DocShift.exe given only
        # file paths opens its window, exactly as for a file dropped on its icon.
        fake = work / "fake.pdf"
        fake.write_text("not a pdf")
        code = run(command, fake, "--output", work / "out")
        if code != 1:
            fail(f"a file that is neither a PDF nor a Word document exited with {code}, not 1")
        print("a file that is neither a PDF nor a Word document is refused, exit code 1")

        # The other direction: a Word document, drawn without Word.
        word = work / "letter.docx"
        make_docx(word)
        code = run(command, word, "--output", work / "out")
        if code != 0:
            fail(f"converting a real Word document exited with {code}")
        written = work / "out" / "letter.pdf"
        if not written.exists():
            fail("it exited cleanly but wrote no PDF")
        with pymupdf.open(str(written)) as converted:
            drawn = " ".join(" ".join(page.get_text().split()) for page in converted)
            pictures = sum(len(page.get_images()) for page in converted)
        for line in ("Shifted by DocShift", "The second paragraph survives", "New York"):
            if line not in drawn:
                fail(f"{line!r} is missing from the PDF")
        if not pictures:
            fail("the picture is missing from the PDF")
        print(f"converted a Word document: its text, its table, and {pictures} picture(s)")
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1:]))
