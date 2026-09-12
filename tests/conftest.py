"""Fixtures shared by the tests.

Qt runs on its offscreen platform, so the window tests need no display: not on
a CI runner, and not on a Windows machine with the screen locked.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


@pytest.fixture
def make_pdf():
    """Write a small PDF with one line of text per page.

    Drawn with PyMuPDF, which pdf2docx cannot run without, so the tests need
    nothing the app does not already install.
    """
    import pymupdf

    def make(path, pages=("A page converted by DocShift",), password=None):
        document = pymupdf.open()
        for text in pages:
            document.new_page().insert_text((72, 72), text, fontsize=14)
        if password:
            document.save(
                str(path),
                encryption=pymupdf.PDF_ENCRYPT_AES_256,
                user_pw=password,
                owner_pw=password + "-owner",
            )
        else:
            document.save(str(path))
        document.close()
        return path

    return make


@pytest.fixture
def make_docx():
    """Write a small Word document, one paragraph per line of text."""
    import docx

    def make(path, paragraphs=("A document converted by DocShift",)):
        document = docx.Document()
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(str(path))
        return path

    return make


@pytest.fixture
def rich_docx():
    """A document holding what the DOCX to PDF engine has to carry across.

    A heading, direct formatting, a nested bullet list, a numbered list, a table
    with a merged cell, a picture and an explicit page break -- plus a header,
    which it cannot carry and has to own up to.
    """
    import docx
    import pymupdf
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Pt, RGBColor

    def make(path):
        document = docx.Document()
        document.add_heading("Quarterly Report", level=1)

        centred = document.add_paragraph()
        centred.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = centred.add_run("Centred and red")
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x33)
        run.font.name = "Georgia"

        line = document.add_paragraph()
        line.add_run("bold").bold = True
        line.add_run(" and ")
        line.add_run("italic").italic = True

        document.add_paragraph("First bullet", style="List Bullet")
        document.add_paragraph("Nested bullet", style="List Bullet 2")
        document.add_paragraph("Step one", style="List Number")

        table = document.add_table(rows=2, cols=3)
        table.style = "Table Grid"
        for column, heading in enumerate(("Region", "Q1", "Q2")):
            table.cell(0, column).text = heading
        table.cell(1, 0).text = "New York"
        table.cell(1, 1).merge(table.cell(1, 2)).text = "one merged cell"

        picture = Path(path).parent / "docshift-test-picture.png"
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40), False)
        pixmap.set_rect(pixmap.irect, (40, 120, 220))
        pixmap.save(str(picture))
        document.add_picture(str(picture), width=Inches(1.5))
        picture.unlink()

        document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("This is on the second page.")
        document.sections[0].header.paragraphs[0].text = "ACME letterhead"

        document.save(str(path))
        return path

    return make


@pytest.fixture
def read_pdf():
    """The text of a PDF, one string per page."""
    import pymupdf

    def read(path) -> list[str]:
        with pymupdf.open(str(path)) as document:
            return [page.get_text() for page in document]

    return read


@pytest.fixture
def read_docx():
    """All the paragraph text in a DOCX, one paragraph per line."""
    import docx

    def read(path) -> str:
        return "\n".join(paragraph.text for paragraph in docx.Document(str(path)).paragraphs)

    return read


@pytest.fixture
def break_pages(monkeypatch):
    """Make pdf2docx fail on the given pages, numbered from 1.

    The same failure a page with layout pdf2docx cannot parse produces, without
    needing a PDF that reliably breaks it.
    """
    from docshift.core.pdf_to_docx import _load_engine

    # Import pdf2docx the way the app does, so its import side effects are
    # contained here too.
    _load_engine()
    from pdf2docx.page.Page import Page

    original = Page.parse

    def break_on(*numbers):
        def parse(page, **settings):
            if page.id + 1 in numbers:
                raise RuntimeError(f"simulated layout failure on page {page.id + 1}")
            return original(page, **settings)

        monkeypatch.setattr(Page, "parse", parse)

    return break_on


@pytest.fixture(scope="session")
def qapp():
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    # The tests close windows on purpose. That must not end the application.
    app.setQuitOnLastWindowClosed(False)
    return app
